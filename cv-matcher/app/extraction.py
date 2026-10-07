import io
from pathlib import Path
import docx
import pymupdf as fitz
from app.config import settings
UPLOAD_DIR = Path(settings.upload_dir).resolve()
FORMATS = {".pdf", ".docx"}

class ExtractionError(Exception):
    pass


def resolve_path(rel_path: str) -> Path:
    path = (UPLOAD_DIR / rel_path).resolve()
    if not path.is_relative_to(UPLOAD_DIR):
        raise ExtractionError("Chemin interdit")
    if not path.is_file():
        raise ExtractionError(f"Fichier introuvable : {rel_path}")
    if path.suffix.lower() not in FORMATS:
        raise ExtractionError(f"Format non supporté : {path.suffix}")
    return path

def extract_pdf(path: Path) -> str:
    with fitz.open(path) as doc:
        return "\n".join(page.get_text() for page in doc)


def extract_docx(path: Path) -> str:
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    # beaucoup de CV Word sont construits avec des tableaux
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)



def ocr_pdf(path: Path) -> str:
    import pytesseract
    from PIL import Image
    
    texts = []
    with fitz.open(path) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=300)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            texts.append(pytesseract.image_to_string(image, lang="fra+eng"))
    return "\n".join(texts)

def extract_text(rel_path: str) -> str:
    path = resolve_path(rel_path)
    if path.suffix.lower() == ".pdf":
        text = extract_pdf(path)
        if len(text.strip()) < 50: # PDF scanné
            text = ocr_pdf(path)
    else:
        text = extract_docx(path)
    if not text.strip():
        raise ExtractionError("Aucun texte lisible dans le document")
    return text